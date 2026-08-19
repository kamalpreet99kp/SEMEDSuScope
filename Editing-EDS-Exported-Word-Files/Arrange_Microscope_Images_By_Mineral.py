"""Build a Word report from manually selected mineral/category folders.

Each selected category folder must contain a ``cropped`` subfolder. The script
does not crop or alter any source image. Run it from PyCharm on Windows, enter
the sample name, add size fractions and category folders, arrange their order,
and choose an output directory. The output filename is the sample name.
"""

from __future__ import annotations

import re
import tempfile
from collections import OrderedDict
from dataclasses import dataclass
from pathlib import Path
from typing import Callable
from tkinter import (
    BOTH,
    END,
    LEFT,
    RIGHT,
    X,
    Button,
    Entry,
    Frame,
    Label,
    Listbox,
    StringVar,
    Tk,
    filedialog,
    messagebox,
    simpledialog,
)
from tkinter import ttk

from PIL import Image, ImageOps
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.shared import Pt


PRESET_FRACTIONS = (
    "5-20/10-20",
    "20-40",
    "40-53",
    "53-75",
    "75-100",
    "100-150/+106",
    "150-300",
    "+300",
)
IMAGE_EXTENSIONS = {".jpg", ".jpeg", ".png", ".tif", ".tiff", ".bmp", ".webp"}
IMAGES_PER_ROW = 3
GAP_POINTS = 5
MARGIN_POINTS = 45
BOTTOM_MARGIN_POINTS = 27
TOP_OFFSET_POINTS = 20
ROWS_PER_PAGE = 6


@dataclass(frozen=True)
class CategorySelection:
    """One category folder assigned to one size fraction."""

    name: str
    folder: Path
    cropped_folder: Path
    images: tuple[Path, ...]


def natural_sort_key(path: Path) -> list[object]:
    """Sort Image2 before Image10 while remaining case-insensitive."""
    return [int(part) if part.isdigit() else part.casefold() for part in re.split(r"(\d+)", path.name)]


def find_cropped_folder(category_folder: Path) -> Path | None:
    """Find a direct child named 'cropped', ignoring capitalization."""
    try:
        return next(
            child
            for child in category_folder.iterdir()
            if child.is_dir() and child.name.casefold() == "cropped"
        )
    except (StopIteration, OSError):
        return None


def find_images(folder: Path) -> tuple[Path, ...]:
    """Return supported images in natural filename order."""
    try:
        images = [
            path
            for path in folder.iterdir()
            if path.is_file() and path.suffix.casefold() in IMAGE_EXTENSIONS
        ]
    except OSError:
        return ()
    return tuple(sorted(images, key=natural_sort_key))


def safe_output_stem(sample_name: str) -> str:
    """Make a Windows-safe filename while retaining the visible sample name."""
    cleaned = re.sub(r'[<>:"/\\|?*]', "_", sample_name).strip().rstrip(".")
    return cleaned or "Microscope Image Report"


def configure_section(section) -> None:
    """Apply the macro margins without adding a Word page header."""
    section.left_margin = Pt(MARGIN_POINTS)
    section.right_margin = Pt(MARGIN_POINTS)
    section.top_margin = Pt(MARGIN_POINTS)
    section.bottom_margin = Pt(BOTTOM_MARGIN_POINTS)


def add_centered_heading(document: Document, text: str, size: int, space_after: int) -> None:
    """Add a centered bold report heading."""
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(space_after)
    run = paragraph.add_run(text)
    run.bold = True
    run.font.size = Pt(size)


def prepare_word_image(image_path: Path, temporary_folder: Path) -> tuple[Path, int, int]:
    """Validate an image, apply EXIF orientation, and provide Word-safe data."""
    with Image.open(image_path) as image:
        requires_conversion = image_path.suffix.casefold() not in {".jpg", ".jpeg", ".png"}
        exif_orientation = image.getexif().get(274, 1)
        orientation_changed = exif_orientation not in (None, 1)

        if not requires_conversion and not orientation_changed:
            return image_path, image.width, image.height

        converted = ImageOps.exif_transpose(image)
        width, height = converted.size
        if converted.mode not in {"RGB", "L"}:
            background = Image.new("RGB", converted.size, "white")
            if "A" in converted.getbands():
                background.paste(converted, mask=converted.getchannel("A"))
            else:
                background.paste(converted.convert("RGB"))
            converted = background
        output_path = temporary_folder / f"{len(list(temporary_folder.iterdir())):06d}.png"
        converted.save(output_path, format="PNG")
        return output_path, width, height


def add_inline_image_block(
    document: Document,
    images: tuple[Path, ...],
    temporary_folder: Path,
    image_processed: Callable[[Path], None] | None = None,
) -> list[str]:
    """Add one editable, automatically reflowing image block for a fraction."""
    section = document.sections[-1]
    page_width = section.page_width.pt
    page_height = section.page_height.pt
    usable_width = page_width - section.left_margin.pt - section.right_margin.pt
    usable_height = page_height - TOP_OFFSET_POINTS - section.top_margin.pt - section.bottom_margin.pt
    # Leave a small tolerance for Word's line-layout rounding. Separator spaces
    # provide approximately the macro's five-point horizontal gap.
    slot_width = (usable_width - 2 * GAP_POINTS - 3) / IMAGES_PER_ROW
    slot_height = (usable_height - (ROWS_PER_PAGE - 1) * GAP_POINTS) / ROWS_PER_PAGE
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(GAP_POINTS)
    paragraph.paragraph_format.line_spacing = Pt(slot_height + GAP_POINTS)
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    errors: list[str] = []

    for image_index, image_path in enumerate(images):
        try:
            prepared_path, pixel_width, pixel_height = prepare_word_image(image_path, temporary_folder)
            scale = min(slot_width / pixel_width, slot_height / pixel_height)
            paragraph.add_run().add_picture(
                str(prepared_path),
                width=Pt(pixel_width * scale),
                height=Pt(pixel_height * scale),
            )
            if image_index < len(images) - 1:
                separator = paragraph.add_run(" ")
                separator.font.size = Pt(GAP_POINTS * 4)
        except (OSError, ValueError) as exc:
            errors.append(f"{image_path}: {exc}")
        finally:
            if image_processed is not None:
                image_processed(image_path)

    return errors


def build_report(
    sample_name: str,
    fraction_order: list[str],
    category_order: list[str],
    selections: OrderedDict[str, list[CategorySelection]],
    output_path: Path,
    progress_callback: Callable[[int, int, str], None] | None = None,
) -> list[str]:
    """Create and save the complete mineral image report."""
    document = Document()
    configure_section(document.sections[0])
    first_category = True
    errors: list[str] = []
    total_images = sum(len(item.images) for group in selections.values() for item in group)
    processed_images = 0

    def image_processed(image_path: Path) -> None:
        nonlocal processed_images
        processed_images += 1
        if progress_callback is not None:
            progress_callback(
                processed_images,
                total_images,
                f"Processing image {processed_images} of {total_images}: {image_path.name}",
            )

    if progress_callback is not None:
        progress_callback(0, total_images, f"Preparing report for {total_images} image(s)...")

    with tempfile.TemporaryDirectory(prefix="mineral_word_report_") as temporary_name:
        temporary_folder = Path(temporary_name)
        for category_key in category_order:
            category_groups: list[tuple[str, str, tuple[Path, ...]]] = []
            for fraction in fraction_order:
                matching = [
                    selection
                    for selection in selections.get(fraction, [])
                    if selection.name.casefold() == category_key
                ]
                if matching:
                    combined_images = tuple(image for selection in matching for image in selection.images)
                    category_groups.append((fraction, matching[0].name, combined_images))

            if not category_groups:
                continue

            category_name = category_groups[0][1]
            if not first_category:
                # An ordinary page break starts every category automatically.
                # Unlike a section break, users can easily delete it in Word
                # if they later want the category to use space on the prior page.
                document.add_page_break()

            if first_category:
                add_centered_heading(document, sample_name, size=18, space_after=6)
            add_centered_heading(document, category_name, size=16, space_after=8)

            for fraction, _name, images in category_groups:
                heading = document.add_paragraph()
                heading.paragraph_format.keep_with_next = True
                heading.paragraph_format.space_before = Pt(6)
                heading.paragraph_format.space_after = Pt(3)
                run = heading.add_run(fraction)
                run.bold = True
                run.font.size = Pt(12)
                errors.extend(add_inline_image_block(document, images, temporary_folder, image_processed))

            first_category = False

        if first_category:
            raise ValueError("No categories with images were available for the report.")
        if progress_callback is not None:
            progress_callback(total_images, total_images, "Saving Word report...")
        document.save(output_path)

    if progress_callback is not None:
        progress_callback(total_images, total_images, "Word report created successfully.")

    return errors


class MineralReportApp:
    """Tkinter interface for assembling fractions and category folders."""

    def __init__(self) -> None:
        self.root = Tk()
        self.root.title("Microscope Images to Word Report")
        self.root.geometry("1050x720")
        self.root.minsize(900, 620)

        self.sample_name = StringVar()
        self.fraction_choice = StringVar(value=PRESET_FRACTIONS[0])
        self.selections: OrderedDict[str, list[CategorySelection]] = OrderedDict()
        self.category_names: OrderedDict[str, str] = OrderedDict()

        self._build_interface()

    def _build_interface(self) -> None:
        top = Frame(self.root)
        top.pack(fill=X, padx=12, pady=(12, 6))
        Label(top, text="Sample name:").pack(side=LEFT)
        Entry(top, textvariable=self.sample_name, width=45).pack(side=LEFT, padx=8)

        chooser = Frame(self.root)
        chooser.pack(fill=X, padx=12, pady=6)
        Label(chooser, text="Size fraction:").pack(side=LEFT)
        self.fraction_combo = ttk.Combobox(
            chooser,
            textvariable=self.fraction_choice,
            values=PRESET_FRACTIONS,
            state="readonly",
            width=24,
        )
        self.fraction_combo.pack(side=LEFT, padx=8)
        Button(chooser, text="Add selected fraction", command=self.add_selected_fraction).pack(side=LEFT, padx=3)
        Button(chooser, text="Add new size fraction", command=self.add_custom_fraction).pack(side=LEFT, padx=3)

        content = Frame(self.root)
        content.pack(fill=BOTH, expand=True, padx=12, pady=8)

        fraction_frame = Frame(content)
        fraction_frame.pack(side=LEFT, fill=BOTH, expand=True, padx=(0, 8))
        Label(fraction_frame, text="Size fractions (report order)", font=("Segoe UI", 10, "bold")).pack(anchor="w")
        self.fraction_list = Listbox(fraction_frame, exportselection=False)
        self.fraction_list.pack(fill=BOTH, expand=True, pady=4)
        self.fraction_list.bind("<<ListboxSelect>>", lambda _event: self.refresh_category_folder_list())
        fraction_buttons = Frame(fraction_frame)
        fraction_buttons.pack(fill=X)
        Button(fraction_buttons, text="Move up", command=lambda: self.move_item(self.fraction_list, -1, "fraction")).pack(side=LEFT)
        Button(fraction_buttons, text="Move down", command=lambda: self.move_item(self.fraction_list, 1, "fraction")).pack(side=LEFT, padx=3)
        Button(fraction_buttons, text="Remove", command=self.remove_fraction).pack(side=RIGHT)

        selected_frame = Frame(content)
        selected_frame.pack(side=LEFT, fill=BOTH, expand=True, padx=8)
        Label(selected_frame, text="Folders for selected fraction", font=("Segoe UI", 10, "bold")).pack(anchor="w")
        self.folder_list = Listbox(selected_frame, exportselection=False)
        self.folder_list.pack(fill=BOTH, expand=True, pady=4)
        folder_buttons = Frame(selected_frame)
        folder_buttons.pack(fill=X)
        Button(folder_buttons, text="Add category folder", command=self.add_category_folder).pack(side=LEFT)
        Button(folder_buttons, text="Add all subfolders", command=self.add_all_category_subfolders).pack(
            side=LEFT, padx=3
        )
        Button(folder_buttons, text="Remove folder", command=self.remove_category_folder).pack(side=RIGHT)

        category_frame = Frame(content)
        category_frame.pack(side=LEFT, fill=BOTH, expand=True, padx=(8, 0))
        Label(category_frame, text="Category order", font=("Segoe UI", 10, "bold")).pack(anchor="w")
        self.category_list = Listbox(category_frame, exportselection=False)
        self.category_list.pack(fill=BOTH, expand=True, pady=4)
        category_buttons = Frame(category_frame)
        category_buttons.pack(fill=X)
        Button(category_buttons, text="Move up", command=lambda: self.move_item(self.category_list, -1, "category")).pack(side=LEFT)
        Button(category_buttons, text="Move down", command=lambda: self.move_item(self.category_list, 1, "category")).pack(side=LEFT, padx=3)

        bottom = Frame(self.root)
        bottom.pack(fill=X, padx=12, pady=(4, 12))
        self.status = Label(bottom, text="Add a size fraction, then add its category folders.", anchor="w")
        self.status.pack(side=LEFT, fill=X, expand=True)
        self.create_button = Button(
            bottom,
            text="Create Word Report",
            command=self.create_report,
            font=("Segoe UI", 10, "bold"),
        )
        self.create_button.pack(side=RIGHT)
        self.progress_bar = ttk.Progressbar(bottom, mode="determinate", length=240)
        self.progress_bar.pack(side=RIGHT, padx=10)

    def selected_fraction(self) -> str | None:
        selection = self.fraction_list.curselection()
        return self.fraction_list.get(selection[0]) if selection else None

    def add_fraction(self, label: str) -> None:
        label = label.strip()
        if not label:
            return
        existing = next((name for name in self.selections if name.casefold() == label.casefold()), None)
        if existing:
            index = list(self.selections).index(existing)
            self.fraction_list.selection_clear(0, END)
            self.fraction_list.selection_set(index)
            self.fraction_list.see(index)
            self.refresh_category_folder_list()
            return
        self.selections[label] = []
        self.fraction_list.insert(END, label)
        index = self.fraction_list.size() - 1
        self.fraction_list.selection_clear(0, END)
        self.fraction_list.selection_set(index)
        self.refresh_category_folder_list()

    def add_selected_fraction(self) -> None:
        self.add_fraction(self.fraction_choice.get())

    def add_custom_fraction(self) -> None:
        label = simpledialog.askstring("New size fraction", "Enter the size-fraction label:", parent=self.root)
        if label is not None:
            self.add_fraction(label)

    def add_category_folder(self) -> None:
        fraction = self.selected_fraction()
        if not fraction:
            messagebox.showwarning("Select a size fraction", "Add or select a size fraction first.", parent=self.root)
            return
        selected = filedialog.askdirectory(title=f"Select a category folder for {fraction}", parent=self.root)
        if not selected:
            return
        folder = Path(selected)
        cropped_folder = find_cropped_folder(folder)
        if cropped_folder is None:
            messagebox.showerror(
                "Cropped folder not found",
                f"The selected category folder does not contain a 'cropped' folder:\n\n{folder}",
                parent=self.root,
            )
            return
        images = find_images(cropped_folder)
        if not images:
            messagebox.showerror(
                "No images found",
                f"No supported images were found in:\n\n{cropped_folder}",
                parent=self.root,
            )
            return
        resolved = folder.resolve()
        if any(item.folder.resolve() == resolved for item in self.selections[fraction]):
            messagebox.showinfo("Already added", "That category folder is already assigned to this fraction.", parent=self.root)
            return
        category_key = folder.name.casefold()
        if any(item.name.casefold() == category_key for item in self.selections[fraction]):
            proceed = messagebox.askyesno(
                "Category already present",
                f"{folder.name} is already present for {fraction}. Add this additional folder to the same category?",
                parent=self.root,
            )
            if not proceed:
                return
        item = CategorySelection(folder.name, folder, cropped_folder, images)
        self.selections[fraction].append(item)
        if category_key not in self.category_names:
            self.category_names[category_key] = folder.name
            self.category_list.insert(END, folder.name)
        self.refresh_category_folder_list()
        self.status.config(text=f"Added {folder.name}: {len(images)} image(s) for {fraction}.")

    def add_all_category_subfolders(self) -> None:
        """Add every valid category beneath one selected parent folder."""
        fraction = self.selected_fraction()
        if not fraction:
            messagebox.showwarning("Select a size fraction", "Add or select a size fraction first.", parent=self.root)
            return
        selected = filedialog.askdirectory(
            title=f"Select the parent folder containing all categories for {fraction}",
            parent=self.root,
        )
        if not selected:
            return

        parent_folder = Path(selected)
        try:
            category_folders = sorted(
                (path for path in parent_folder.iterdir() if path.is_dir()),
                key=natural_sort_key,
            )
        except OSError as exc:
            messagebox.showerror("Folder could not be read", str(exc), parent=self.root)
            return

        added_categories = 0
        added_images = 0
        skipped: list[str] = []
        existing_paths = {item.folder.resolve() for item in self.selections[fraction]}

        for folder in category_folders:
            cropped_folder = find_cropped_folder(folder)
            if cropped_folder is None:
                skipped.append(f"{folder.name}: no cropped folder")
                continue
            images = find_images(cropped_folder)
            if not images:
                skipped.append(f"{folder.name}: cropped folder has no supported images")
                continue
            resolved = folder.resolve()
            if resolved in existing_paths:
                skipped.append(f"{folder.name}: already added")
                continue

            item = CategorySelection(folder.name, folder, cropped_folder, images)
            self.selections[fraction].append(item)
            existing_paths.add(resolved)
            category_key = folder.name.casefold()
            if category_key not in self.category_names:
                self.category_names[category_key] = folder.name
                self.category_list.insert(END, folder.name)
            added_categories += 1
            added_images += len(images)

        self.refresh_category_folder_list()
        self.status.config(
            text=f"Added {added_categories} categories and {added_images} images for {fraction}."
        )
        summary = (
            f"Parent folder: {parent_folder}\n\n"
            f"Categories added: {added_categories}\n"
            f"Images added: {added_images}\n"
            f"Folders skipped: {len(skipped)}"
        )
        if skipped:
            displayed = skipped[:20]
            summary += "\n\nSkipped:\n" + "\n".join(displayed)
            if len(skipped) > len(displayed):
                summary += f"\n...and {len(skipped) - len(displayed)} more"
        messagebox.showinfo("Category folders processed", summary, parent=self.root)

    def refresh_category_folder_list(self) -> None:
        self.folder_list.delete(0, END)
        fraction = self.selected_fraction()
        if not fraction:
            return
        for item in self.selections[fraction]:
            self.folder_list.insert(END, f"{item.name} — {len(item.images)} images — {item.folder}")

    def remove_category_folder(self) -> None:
        fraction = self.selected_fraction()
        selected = self.folder_list.curselection()
        if not fraction or not selected:
            return
        del self.selections[fraction][selected[0]]
        self.rebuild_category_order()
        self.refresh_category_folder_list()

    def remove_fraction(self) -> None:
        selected = self.fraction_list.curselection()
        if not selected:
            return
        index = selected[0]
        fraction = self.fraction_list.get(index)
        if self.selections[fraction] and not messagebox.askyesno(
            "Remove size fraction",
            f"Remove {fraction} and all of its selected category folders?",
            parent=self.root,
        ):
            return
        del self.selections[fraction]
        self.fraction_list.delete(index)
        self.rebuild_category_order()
        self.refresh_category_folder_list()

    def rebuild_category_order(self) -> None:
        used = {
            selection.name.casefold(): selection.name
            for group in self.selections.values()
            for selection in group
        }
        old_order = [self.category_list.get(i).casefold() for i in range(self.category_list.size())]
        new_order = [key for key in old_order if key in used]
        new_order.extend(key for key in used if key not in new_order)
        self.category_names = OrderedDict((key, used[key]) for key in new_order)
        self.category_list.delete(0, END)
        for name in self.category_names.values():
            self.category_list.insert(END, name)

    def move_item(self, listbox: Listbox, direction: int, item_type: str) -> None:
        selected = listbox.curselection()
        if not selected:
            return
        index = selected[0]
        destination = index + direction
        if destination < 0 or destination >= listbox.size():
            return
        text = listbox.get(index)
        listbox.delete(index)
        listbox.insert(destination, text)
        listbox.selection_set(destination)
        listbox.see(destination)

        if item_type == "fraction":
            reordered = OrderedDict()
            for position in range(listbox.size()):
                label = listbox.get(position)
                reordered[label] = self.selections[label]
            self.selections = reordered
            self.refresh_category_folder_list()
        else:
            keys = [listbox.get(position).casefold() for position in range(listbox.size())]
            self.category_names = OrderedDict((key, self.category_names[key]) for key in keys)

    def update_progress(self, completed: int, total: int, message: str) -> None:
        """Update report-generation status while the document is being built."""
        self.progress_bar.configure(maximum=max(total, 1), value=completed)
        self.status.config(text=message)
        self.root.update_idletasks()

    def create_report(self) -> None:
        sample_name = self.sample_name.get().strip()
        if not sample_name:
            messagebox.showwarning("Sample name required", "Enter the sample name first.", parent=self.root)
            return
        if not any(self.selections.values()):
            messagebox.showwarning("No category folders", "Add at least one category folder first.", parent=self.root)
            return
        output_directory = filedialog.askdirectory(title="Select the output folder", parent=self.root)
        if not output_directory:
            return
        output_path = Path(output_directory) / f"{safe_output_stem(sample_name)}.docx"
        if output_path.exists() and not messagebox.askyesno(
            "Replace existing report?",
            f"This file already exists:\n\n{output_path}\n\nReplace it?",
            parent=self.root,
        ):
            return

        fraction_order = [self.fraction_list.get(i) for i in range(self.fraction_list.size())]
        category_order = [self.category_list.get(i).casefold() for i in range(self.category_list.size())]
        self.create_button.config(state="disabled")
        self.progress_bar.configure(maximum=1, value=0)
        try:
            errors = build_report(
                sample_name,
                fraction_order,
                category_order,
                self.selections,
                output_path,
                progress_callback=self.update_progress,
            )
        except (OSError, ValueError) as exc:
            messagebox.showerror("Report could not be created", str(exc), parent=self.root)
            self.status.config(text="Report creation failed.")
            return
        finally:
            self.create_button.config(state="normal")

        total_images = sum(len(item.images) for group in self.selections.values() for item in group)
        message = f"Report saved successfully.\n\n{output_path}\n\nImages selected: {total_images}"
        if errors:
            message += f"\nImages skipped because they could not be read: {len(errors)}"
            error_log = output_path.with_suffix(".errors.txt")
            error_log.write_text("\n".join(errors), encoding="utf-8")
            message += f"\nError details: {error_log}"
        messagebox.showinfo("Word report created", message, parent=self.root)
        self.status.config(text=f"Saved: {output_path}")

    def run(self) -> None:
        self.root.mainloop()


if __name__ == "__main__":
    MineralReportApp().run()
