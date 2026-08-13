"""GUI calculator for relocating SEM stage positions after remounting a sample.

The same two physical reference points (A and B) are measured before and after
the sample is remounted.  A rigid 2-D transform is then applied to every old
position of interest.  The program never changes an imported file; results are
saved to a new Excel or CSV file selected by the user.
"""

from __future__ import annotations

import csv
import math
import os
import tkinter as tk
from dataclasses import dataclass
from tkinter import filedialog, messagebox, ttk


@dataclass(frozen=True)
class Point:
    x: float
    y: float


@dataclass(frozen=True)
class Transform:
    old_a: Point
    new_a: Point
    cosine: float
    sine: float
    scale: float
    rotation_degrees: float
    old_reference_distance: float
    new_reference_distance: float

    def apply(self, point: Point) -> Point:
        """Return the new coordinates for an old coordinate."""
        relative_x = point.x - self.old_a.x
        relative_y = point.y - self.old_a.y
        rotated_x = relative_x * self.cosine - relative_y * self.sine
        rotated_y = relative_x * self.sine + relative_y * self.cosine
        return Point(
            self.new_a.x + self.scale * rotated_x,
            self.new_a.y + self.scale * rotated_y,
        )


def calculate_transform(
    old_a: Point,
    old_b: Point,
    new_a: Point,
    new_b: Point,
    use_scale: bool = False,
) -> Transform:
    """Build a translation/rotation transform from two reference-point pairs."""
    old_dx = old_b.x - old_a.x
    old_dy = old_b.y - old_a.y
    new_dx = new_b.x - new_a.x
    new_dy = new_b.y - new_a.y
    old_distance = math.hypot(old_dx, old_dy)
    new_distance = math.hypot(new_dx, new_dy)

    if old_distance == 0:
        raise ValueError("Old reference points A and B cannot be the same position.")
    if new_distance == 0:
        raise ValueError("New reference points A and B cannot be the same position.")

    old_angle = math.atan2(old_dy, old_dx)
    new_angle = math.atan2(new_dy, new_dx)
    rotation = new_angle - old_angle
    scale = new_distance / old_distance if use_scale else 1.0
    return Transform(
        old_a=old_a,
        new_a=new_a,
        cosine=math.cos(rotation),
        sine=math.sin(rotation),
        scale=scale,
        rotation_degrees=math.degrees(rotation),
        old_reference_distance=old_distance,
        new_reference_distance=new_distance,
    )


class RelocationCalculator(tk.Tk):
    """Editable Windows-friendly workflow for SEM coordinate relocation."""

    def __init__(self) -> None:
        super().__init__()
        self.title("SEM Position Relocation Calculator")
        self.geometry("1020x700")
        self.minsize(850, 600)
        self.reference_entries: dict[str, ttk.Entry] = {}
        self.rows: list[dict[str, object]] = []
        self.use_scale = tk.BooleanVar(value=False)
        self.status = tk.StringVar(value="Enter reference positions and positions of interest.")
        self._build_interface()
        self.add_row()

    def _build_interface(self) -> None:
        heading = ttk.Label(
            self,
            text="Relocate SEM Positions After Sample Remounting",
            font=("Segoe UI", 15, "bold"),
        )
        heading.pack(pady=(12, 4))
        ttk.Label(
            self,
            text=("Enter the old and newly measured SEM coordinates for the same marks A and B, "
                  "then enter all old positions of interest."),
        ).pack(pady=(0, 10))

        references = ttk.LabelFrame(self, text="1. Reference positions")
        references.pack(fill="x", padx=12, pady=4)
        for column, text in enumerate(("Reference", "Old X", "Old Y", "New X", "New Y")):
            ttk.Label(references, text=text, font=("Segoe UI", 9, "bold")).grid(
                row=0, column=column, padx=7, pady=5
            )
        for row_number, mark in enumerate(("A", "B"), start=1):
            ttk.Label(references, text=mark).grid(row=row_number, column=0, padx=7, pady=4)
            for column, coordinate in enumerate(("old_x", "old_y", "new_x", "new_y"), start=1):
                entry = ttk.Entry(references, width=19)
                entry.grid(row=row_number, column=column, padx=7, pady=4, sticky="ew")
                self.reference_entries[f"{mark.lower()}_{coordinate}"] = entry
            references.columnconfigure(column, weight=1)

        option_frame = ttk.Frame(references)
        option_frame.grid(row=3, column=0, columnspan=5, sticky="w", padx=7, pady=6)
        ttk.Checkbutton(
            option_frame,
            text="Apply measured scale change (normally leave OFF for the same SEM stage)",
            variable=self.use_scale,
        ).pack(side="left")

        positions = ttk.LabelFrame(self, text="2. Positions of interest (editable)")
        positions.pack(fill="both", expand=True, padx=12, pady=8)
        toolbar = ttk.Frame(positions)
        toolbar.pack(fill="x", padx=6, pady=5)
        ttk.Button(toolbar, text="Add position", command=self.add_row).pack(side="left", padx=3)
        ttk.Button(toolbar, text="Remove selected", command=self.remove_selected).pack(side="left", padx=3)
        ttk.Button(toolbar, text="Import CSV / Excel", command=self.import_positions).pack(side="left", padx=3)
        ttk.Button(toolbar, text="Clear all", command=self.clear_rows).pack(side="left", padx=3)

        table_holder = ttk.Frame(positions)
        table_holder.pack(fill="both", expand=True, padx=6, pady=(0, 6))
        self.canvas = tk.Canvas(table_holder, highlightthickness=0)
        scrollbar = ttk.Scrollbar(table_holder, orient="vertical", command=self.canvas.yview)
        self.row_frame = ttk.Frame(self.canvas)
        self.row_frame.bind(
            "<Configure>", lambda _event: self.canvas.configure(scrollregion=self.canvas.bbox("all"))
        )
        self.canvas.create_window((0, 0), window=self.row_frame, anchor="nw")
        self.canvas.configure(yscrollcommand=scrollbar.set)
        self.canvas.pack(side="left", fill="both", expand=True)
        scrollbar.pack(side="right", fill="y")
        for column, label in enumerate(("Select", "ID / Description", "Old X", "Old Y", "Calculated New X", "Calculated New Y")):
            ttk.Label(self.row_frame, text=label, font=("Segoe UI", 9, "bold")).grid(
                row=0, column=column, padx=5, pady=4, sticky="ew"
            )
            self.row_frame.columnconfigure(column, weight=1 if column else 0)

        actions = ttk.Frame(self)
        actions.pack(fill="x", padx=12, pady=(0, 8))
        ttk.Button(actions, text="3. Calculate / Recalculate", command=self.calculate).pack(side="left", padx=3)
        ttk.Button(actions, text="4. Save results", command=self.save_results).pack(side="left", padx=3)
        ttk.Label(actions, textvariable=self.status).pack(side="left", padx=12)

    def add_row(self, identifier: str = "", old_x: object = "", old_y: object = "") -> None:
        row_index = len(self.rows) + 1
        selected = tk.BooleanVar(value=False)
        widgets: list[tk.Widget] = [ttk.Checkbutton(self.row_frame, variable=selected)]
        values = (identifier, old_x, old_y, "", "")
        entries: list[ttk.Entry] = []
        for value in values:
            entry = ttk.Entry(self.row_frame, width=22)
            entry.insert(0, str(value))
            entries.append(entry)
            widgets.append(entry)
        for column, widget in enumerate(widgets):
            widget.grid(row=row_index, column=column, padx=5, pady=3, sticky="ew")
        self.rows.append({"selected": selected, "widgets": widgets, "entries": entries})

    def _rebuild_rows(self, retained_rows: list[tuple[str, str, str]]) -> None:
        for row in self.rows:
            for widget in row["widgets"]:  # type: ignore[union-attr]
                widget.destroy()
        self.rows.clear()
        for values in retained_rows:
            self.add_row(*values)
        if not self.rows:
            self.add_row()

    def remove_selected(self) -> None:
        retained = []
        for row in self.rows:
            if not row["selected"].get():  # type: ignore[union-attr]
                entries = row["entries"]  # type: ignore[assignment]
                retained.append(tuple(entry.get() for entry in entries[:3]))
        self._rebuild_rows(retained)

    def clear_rows(self) -> None:
        if messagebox.askyesno("Clear positions", "Remove all positions of interest?"):
            self._rebuild_rows([])

    @staticmethod
    def _find_column(columns: list[str], choices: tuple[str, ...]) -> str | None:
        normalized = {column.strip().lower(): column for column in columns}
        return next((normalized[name] for name in choices if name in normalized), None)

    def import_positions(self) -> None:
        path = filedialog.askopenfilename(
            title="Select positions file",
            filetypes=[("Supported files", "*.csv *.xlsx"), ("CSV", "*.csv"), ("Excel", "*.xlsx")],
        )
        if not path:
            return
        try:
            if path.lower().endswith(".csv"):
                with open(path, newline="", encoding="utf-8-sig") as file:
                    records = list(csv.DictReader(file))
            else:
                from openpyxl import load_workbook

                workbook = load_workbook(path, read_only=True, data_only=True)
                sheet = workbook.active
                data = list(sheet.iter_rows(values_only=True))
                if not data:
                    raise ValueError("The selected worksheet is empty.")
                headers = [str(value or "") for value in data[0]]
                records = [dict(zip(headers, values)) for values in data[1:]]
            if not records:
                raise ValueError("No data rows were found.")
            columns = list(records[0])
            x_column = self._find_column(columns, ("old x", "old_x", "x", "stage x", "stage x (mm)"))
            y_column = self._find_column(columns, ("old y", "old_y", "y", "stage y", "stage y (mm)"))
            id_column = self._find_column(columns, ("id", "identifier", "grain", "grain id", "description", "class"))
            if not x_column or not y_column:
                raise ValueError("Could not find X and Y columns. Use headings 'Old X' and 'Old Y'.")
            imported = []
            for number, record in enumerate(records, start=1):
                if record.get(x_column) in (None, "") and record.get(y_column) in (None, ""):
                    continue
                imported.append((record.get(id_column, number) if id_column else number,
                                 record.get(x_column, ""), record.get(y_column, "")))
            self._rebuild_rows([(str(a), str(b), str(c)) for a, b, c in imported])
            self.status.set(f"Imported {len(imported)} positions. You can edit or add rows before calculating.")
        except Exception as error:
            messagebox.showerror("Import failed", str(error))

    @staticmethod
    def _number(entry: ttk.Entry, label: str) -> float:
        text = entry.get().strip().replace(",", ".")
        try:
            return float(text)
        except ValueError as error:
            raise ValueError(f"{label} must be a number.") from error

    def _get_transform(self) -> Transform:
        value = lambda key: self._number(self.reference_entries[key], key.replace("_", " ").title())
        return calculate_transform(
            Point(value("a_old_x"), value("a_old_y")),
            Point(value("b_old_x"), value("b_old_y")),
            Point(value("a_new_x"), value("a_new_y")),
            Point(value("b_new_x"), value("b_new_y")),
            self.use_scale.get(),
        )

    def calculate(self) -> None:
        try:
            transform = self._get_transform()
            calculated = 0
            for row_number, row in enumerate(self.rows, start=1):
                entries = row["entries"]  # type: ignore[assignment]
                if not entries[1].get().strip() and not entries[2].get().strip():
                    entries[3].delete(0, tk.END)
                    entries[4].delete(0, tk.END)
                    continue
                result = transform.apply(Point(
                    self._number(entries[1], f"Row {row_number} Old X"),
                    self._number(entries[2], f"Row {row_number} Old Y"),
                ))
                for entry, value in zip(entries[3:5], (result.x, result.y)):
                    entry.delete(0, tk.END)
                    entry.insert(0, f"{value:.6f}")
                calculated += 1
            distance_difference = transform.new_reference_distance - transform.old_reference_distance
            self.status.set(
                f"Calculated {calculated} positions | Rotation: {transform.rotation_degrees:.6f}° | "
                f"A–B distance difference: {distance_difference:.6f} | Scale used: {transform.scale:.9f}"
            )
        except ValueError as error:
            messagebox.showerror("Check the entered values", str(error))

    def _output_rows(self) -> list[list[object]]:
        transform = self._get_transform()
        output = []
        for row_number, row in enumerate(self.rows, start=1):
            entries = row["entries"]  # type: ignore[assignment]
            if not entries[1].get().strip() and not entries[2].get().strip():
                continue
            old = Point(self._number(entries[1], f"Row {row_number} Old X"),
                        self._number(entries[2], f"Row {row_number} Old Y"))
            new = transform.apply(old)
            output.append([entries[0].get().strip() or row_number, old.x, old.y, new.x, new.y])
        return output

    def save_results(self) -> None:
        try:
            self.calculate()
            rows = self._output_rows()
            if not rows:
                raise ValueError("Enter at least one position of interest before saving.")
            path = filedialog.asksaveasfilename(
                title="Save calculated positions as a new file",
                defaultextension=".xlsx",
                initialfile="Calculated_New_SEM_Positions.xlsx",
                filetypes=[("Excel workbook", "*.xlsx"), ("CSV", "*.csv")],
            )
            if not path:
                return
            headers = ["ID / Description", "Old SEM X", "Old SEM Y", "Calculated New SEM X", "Calculated New SEM Y"]
            if path.lower().endswith(".csv"):
                with open(path, "w", newline="", encoding="utf-8-sig") as file:
                    writer = csv.writer(file)
                    writer.writerow(headers)
                    writer.writerows(rows)
            else:
                from openpyxl import Workbook

                workbook = Workbook()
                results = workbook.active
                results.title = "Calculated Positions"
                results.append(headers)
                for row in rows:
                    results.append(row)
                inputs = workbook.create_sheet("Reference Inputs")
                inputs.append(["Reference", "Old X", "Old Y", "New X", "New Y"])
                for mark in ("a", "b"):
                    inputs.append([mark.upper()] + [self.reference_entries[f"{mark}_{name}"].get()
                                                   for name in ("old_x", "old_y", "new_x", "new_y")])
                inputs.append([])
                inputs.append(["Scale correction applied", "Yes" if self.use_scale.get() else "No"])
                for sheet in workbook.worksheets:
                    sheet.freeze_panes = "A2"
                    for column in sheet.columns:
                        letter = column[0].column_letter
                        sheet.column_dimensions[letter].width = min(35, max(12, max(len(str(cell.value or "")) for cell in column) + 2))
                workbook.save(path)
            messagebox.showinfo("Results saved", f"A new output file was created:\n{os.path.normpath(path)}")
        except (ValueError, OSError) as error:
            messagebox.showerror("Could not save results", str(error))


def main() -> None:
    app = RelocationCalculator()
    app.mainloop()


if __name__ == "__main__":
    main()
